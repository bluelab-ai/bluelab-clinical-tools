#!/usr/bin/env python3
"""Generate the two-group time-to-event shell as one Word table."""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old_borders = tc_pr.find(qn("w:tcBorders"))
    if old_borders is not None:
        tc_pr.remove(old_borders)

    borders = OxmlElement("w:tcBorders")
    for side, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        border = OxmlElement(f"w:{side}")
        if spec is None or spec == "nil":
            border.set(qn("w:val"), "nil")
        else:
            border.set(qn("w:val"), "single")
            border.set(qn("w:color"), spec.get("color", "auto"))
            border.set(qn("w:sz"), str(spec.get("sz", 4)))
            border.set(qn("w:space"), "0")
        borders.append(border)
    tc_pr.append(borders)


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)

    run = paragraph.add_run(str(text))
    run.font.name = "Times New Roman"
    run.font.size = Pt(10.5)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "宋体")
    r_fonts.set(qn("w:cs"), "Times New Roman")


def format_header(label: str) -> str:
    return label.replace(" (N=n1)", "\n(N=n1)")


def apply_three_line_borders(table) -> None:
    last_row = len(table.rows) - 1
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            if row_index == 0:
                set_cell_borders(
                    cell,
                    top={"sz": 4, "color": "auto"},
                    bottom={"sz": 4, "color": "000000"},
                    left="nil",
                    right="nil",
                )
            elif row_index == last_row:
                set_cell_borders(
                    cell,
                    top="nil",
                    bottom={"sz": 4, "color": "auto"},
                    left="nil",
                    right="nil",
                )
            else:
                set_cell_borders(cell, top="nil", bottom="nil", left="nil", right="nil")


def set_table_width(table) -> None:
    table.autofit = True
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.insert(0, table_width)
    table_width.set(qn("w:w"), "5000")
    table_width.set(qn("w:type"), "pct")


def generate_docx(filled_data: dict, output_path) -> None:
    """Render all endpoints into one table because merge keeps only the first table."""
    document = Document()
    section = document.sections[0]
    section.left_margin = Cm(1.91)
    section.right_margin = Cm(1.91)

    columns = filled_data.get("columns", [])
    rows = filled_data.get("rows", [])
    table = document.add_table(rows=1 + len(rows), cols=len(columns))
    set_table_width(table)

    for index, column in enumerate(columns):
        set_cell_text(table.cell(0, index), format_header(column["name"]))

    for row_index, row in enumerate(rows, start=1):
        set_cell_text(table.cell(row_index, 0), row.get("indicator", ""))
        set_cell_text(table.cell(row_index, 1), row.get("metric", ""))
        values = row.get("values", ["", ""])
        set_cell_text(table.cell(row_index, 2), values[0] if values else "")
        set_cell_text(table.cell(row_index, 3), values[1] if len(values) > 1 else "")

    header_properties = table.rows[0]._tr.get_or_add_trPr()
    header_properties.append(OxmlElement("w:tblHeader"))
    apply_three_line_borders(table)
    document.save(output_path)
